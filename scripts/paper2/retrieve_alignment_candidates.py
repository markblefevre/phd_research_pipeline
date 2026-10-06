#!/usr/bin/env python3
"""
Compare three retrieval-context approaches for Japanese MD&A information novelty.

For each benchmark CURRENT sentence and each embedding model:

Approach 1
----------
Retrieve the PRIOR sentence with maximum cosine similarity.

Approach 2
----------
Take that exact same best prior sentence and expand it by +/- one prior sentence.

Approach 3
----------
Take that exact same best prior sentence and return the entire PRIOR paragraph
that contains it.

Important:
- Approaches 2 and 3 DO NOT perform a second retrieval.
- They only change how much prior-year context is supplied around the best
  sentence found in Approach 1.
- This script tests retrieval/context construction only. It does not classify
  persistent versus novel information.
- Paragraphs are defined by blank-line boundaries in the source text
  (two or more newlines, allowing intervening whitespace).

Example
-------
python scripts/paper2/retrieve_alignment_candidates.py \
    --benchmark japanese_mdna_alignment_candidates.xlsx \
    --output japanese_mdna_retrieval_context_test.xlsx \
    --top-k 3
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd


MODELS = {
    "ruri": "cl-nagoya/ruri-v3-310m",
    "sarashina": "sbintuitions/sarashina-embedding-v2-1b",
}

MIN_SENTENCE_CHARS = 12


def read_text(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError(path)
    return path.read_text(encoding="utf-8-sig", errors="replace")


def normalize_sentence(text: str) -> str:
    text = str(text).replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def split_japanese_sentences(
    text: str,
    min_chars: int = MIN_SENTENCE_CHARS,
) -> list[str]:
    """
    Match the conservative sentence splitting used in the benchmark builder.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    chunks = re.split(r"(?<=[。！？!?])|\n+", text)
    sentences = [normalize_sentence(x) for x in chunks]
    return [x for x in sentences if len(x) >= min_chars]


def split_paragraphs(text: str) -> list[str]:
    """
    Split source text on blank lines while preserving single newlines
    inside each paragraph.

    This is important because the benchmark sentence splitter treats
    single newlines as sentence boundaries.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    # Two or more newline-separated lines, allowing whitespace on blank lines.
    parts = re.split(r"\n[ \t]*\n+", text)

    paragraphs = []
    for part in parts:
        part = part.strip()
        if part:
            paragraphs.append(part)

    return paragraphs


def map_sentences_to_paragraphs(
    text: str,
    min_chars: int = MIN_SENTENCE_CHARS,
) -> tuple[list[str], list[str], list[int]]:
    """
    Reconstruct the sentence list paragraph by paragraph and map each sentence
    index to its paragraph index.

    Returns:
        sentences
        paragraphs
        sentence_to_paragraph
    """
    paragraphs = split_paragraphs(text)

    sentences: list[str] = []
    sentence_to_paragraph: list[int] = []

    for p_idx, paragraph in enumerate(paragraphs):
        p_sentences = split_japanese_sentences(
            paragraph,
            min_chars=min_chars,
        )
        for sent in p_sentences:
            sentences.append(sent)
            sentence_to_paragraph.append(p_idx)

    return sentences, paragraphs, sentence_to_paragraph


def choose_device() -> str:
    import torch

    if torch.cuda.is_available():
        return "cuda"
    if getattr(torch.backends, "mps", None) and torch.backends.mps.is_available():
        return "mps"
    return "cpu"


def encode(
    model,
    texts: list[str],
    batch_size: int,
) -> np.ndarray:
    return np.asarray(
        model.encode(
            texts,
            batch_size=batch_size,
            normalize_embeddings=True,
            show_progress_bar=True,
        )
    )


def top_k_matches(
    query_embeddings: np.ndarray,
    candidate_embeddings: np.ndarray,
    k: int,
) -> tuple[np.ndarray, np.ndarray]:
    """
    Return top-k prior-sentence indices and cosine similarities.

    Embeddings are normalized, so dot product equals cosine similarity.
    """
    k = min(k, candidate_embeddings.shape[0])
    scores = query_embeddings @ candidate_embeddings.T

    part = np.argpartition(-scores, kth=k - 1, axis=1)[:, :k]
    part_scores = np.take_along_axis(scores, part, axis=1)

    order = np.argsort(-part_scores, axis=1)
    top_idx = np.take_along_axis(part, order, axis=1)
    top_scores = np.take_along_axis(part_scores, order, axis=1)

    return top_idx, top_scores


def local_window(
    sentences: list[str],
    center: int,
    radius: int = 1,
) -> tuple[int, int, str]:
    start = max(0, center - radius)
    end = min(len(sentences) - 1, center + radius)
    text = " ".join(sentences[start : end + 1])
    return start, end, text


def paragraph_for_sentence(
    paragraphs: list[str],
    sentence_to_paragraph: list[int],
    sentence_index: int,
) -> tuple[int, str]:
    p_idx = sentence_to_paragraph[sentence_index]
    return p_idx, paragraphs[p_idx]


def parse_optional_int(value):
    if pd.isna(value) or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def add_model_results(
    frame: pd.DataFrame,
    short_name: str,
    model_name: str,
    prior_sentences: list[str],
    prior_paragraphs: list[str],
    sentence_to_paragraph: list[int],
    current_sentences: list[str],
    batch_size: int,
    top_k: int,
    device: str,
) -> pd.DataFrame:
    from sentence_transformers import SentenceTransformer

    print(f"\nLoading {short_name}: {model_name}")
    print(f"Device: {device}")

    model = SentenceTransformer(model_name, device=device)

    print(f"Encoding {len(prior_sentences):,} prior sentences...")
    prior_e = encode(model, prior_sentences, batch_size)

    print(f"Encoding {len(current_sentences):,} benchmark current sentences...")
    current_e = encode(model, current_sentences, batch_size)

    idx, score = top_k_matches(current_e, prior_e, top_k)

    best_indices = idx[:, 0].astype(int)

    # ------------------------------------------------------------------
    # Approach 1: best sentence by max cosine similarity
    # ------------------------------------------------------------------
    frame[f"{short_name}_a1_prior_index"] = best_indices
    frame[f"{short_name}_a1_prior_sentence"] = [
        prior_sentences[i] for i in best_indices
    ]
    frame[f"{short_name}_a1_cosine"] = score[:, 0]

    # ------------------------------------------------------------------
    # Approach 2: same best sentence +/- 1 sentence
    # ------------------------------------------------------------------
    a2 = [
        local_window(prior_sentences, int(i), radius=1)
        for i in best_indices
    ]

    frame[f"{short_name}_a2_start_index"] = [x[0] for x in a2]
    frame[f"{short_name}_a2_end_index"] = [x[1] for x in a2]
    frame[f"{short_name}_a2_context"] = [x[2] for x in a2]

    # ------------------------------------------------------------------
    # Approach 3: entire paragraph containing same best sentence
    # ------------------------------------------------------------------
    a3 = [
        paragraph_for_sentence(
            prior_paragraphs,
            sentence_to_paragraph,
            int(i),
        )
        for i in best_indices
    ]

    frame[f"{short_name}_a3_paragraph_index"] = [x[0] for x in a3]
    frame[f"{short_name}_a3_paragraph"] = [x[1] for x in a3]
    frame[f"{short_name}_a3_paragraph_chars"] = [
        len(x[1]) for x in a3
    ]

    # Paragraph sentence count, useful for spotting giant paragraphs.
    paragraph_sentence_counts = [
        len(split_japanese_sentences(x[1]))
        for x in a3
    ]
    frame[f"{short_name}_a3_paragraph_sentence_count"] = (
        paragraph_sentence_counts
    )

    # Retain compact top-k sentence retrieval diagnostics.
    frame[f"{short_name}_topk_prior_indices"] = [
        ",".join(str(int(x)) for x in row) for row in idx
    ]
    frame[f"{short_name}_topk_cosines"] = [
        ",".join(f"{float(x):.6f}" for x in row)
        for row in score
    ]

    # If human gold alignment exists, evaluate Approach 1 retrieval.
    # Approaches 2/3 necessarily contain the Approach-1 sentence, so their
    # substantive value should be judged by whether they add useful evidence.
    if "true_prior_sentence_index" in frame.columns:
        gold = [
            parse_optional_int(x)
            for x in frame["true_prior_sentence_index"].tolist()
        ]

        frame[f"{short_name}_a1_gold_exact"] = [
            (g is not None and int(s) == g)
            for g, s in zip(gold, best_indices)
        ]

        frame[f"{short_name}_a1_gold_topk"] = [
            (g is not None and g in set(map(int, row)))
            for g, row in zip(gold, idx)
        ]

        frame[f"{short_name}_a2_contains_gold"] = [
            (
                g is not None
                and start <= g <= end
            )
            for g, (start, end, _) in zip(gold, a2)
        ]

        frame[f"{short_name}_a3_contains_gold"] = [
            (
                g is not None
                and sentence_to_paragraph[g] == p_idx
            )
            if g is not None and 0 <= g < len(sentence_to_paragraph)
            else False
            for g, (p_idx, _) in zip(gold, a3)
        ]

    del model
    return frame


def build_summary(
    annotation: pd.DataFrame,
    company_stats: list[dict],
    device: str,
    top_k: int,
) -> pd.DataFrame:
    rows = [
        ["benchmark_rows", len(annotation)],
        ["models", ", ".join(MODELS.keys())],
        ["device", device],
        ["top_k_sentence_candidates_saved", top_k],
        ["approach_1", "best prior sentence by maximum cosine similarity"],
        ["approach_2", "Approach-1 sentence +/- one sentence"],
        ["approach_3", "entire prior paragraph containing Approach-1 sentence"],
        ["paragraph_boundary", "blank line: newline + optional whitespace + newline"],
    ]

    for s in company_stats:
        company = s["company"]
        rows.extend(
            [
                [f"{company}_prior_sentences", s["prior_sentences"]],
                [f"{company}_prior_paragraphs", s["prior_paragraphs"]],
                [
                    f"{company}_median_prior_paragraph_sentences",
                    s["median_paragraph_sentences"],
                ],
                [
                    f"{company}_max_prior_paragraph_sentences",
                    s["max_paragraph_sentences"],
                ],
            ]
        )

    for short in MODELS:
        rows.extend(
            [
                [
                    f"{short}_mean_a1_cosine",
                    float(annotation[f"{short}_a1_cosine"].mean()),
                ],
                [
                    f"{short}_median_a3_paragraph_chars",
                    float(
                        annotation[
                            f"{short}_a3_paragraph_chars"
                        ].median()
                    ),
                ],
                [
                    f"{short}_median_a3_paragraph_sentence_count",
                    float(
                        annotation[
                            f"{short}_a3_paragraph_sentence_count"
                        ].median()
                    ),
                ],
            ]
        )

        if "true_prior_sentence_index" in annotation.columns:
            gold_mask = annotation["true_prior_sentence_index"].apply(
                lambda x: parse_optional_int(x) is not None
            )
            n_gold = int(gold_mask.sum())
            rows.append([f"{short}_gold_rows", n_gold])

            if n_gold:
                for metric in [
                    "a1_gold_exact",
                    "a1_gold_topk",
                    "a2_contains_gold",
                    "a3_contains_gold",
                ]:
                    rows.append(
                        [
                            f"{short}_{metric}_rate",
                            float(
                                annotation.loc[
                                    gold_mask,
                                    f"{short}_{metric}",
                                ].mean()
                            ),
                        ]
                    )

    return pd.DataFrame(rows, columns=["metric", "value"])


def write_workbook(
    annotation: pd.DataFrame,
    summary: pd.DataFrame,
    metadata: pd.DataFrame,
    output: Path,
) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    output.parent.mkdir(parents=True, exist_ok=True)

    guide = pd.DataFrame(
        [
            [
                "Approach 1",
                "Retrieve the single prior-year sentence with maximum embedding cosine similarity.",
            ],
            [
                "Approach 2",
                "Use the exact same Approach-1 sentence but add one prior sentence before and one after it.",
            ],
            [
                "Approach 3",
                "Use the entire prior-year paragraph containing the exact same Approach-1 sentence.",
            ],
            [
                "Paragraph boundary",
                r"Blank line in the source MD&A: \n followed by optional whitespace and another \n.",
            ],
            [
                "Key design",
                "Approaches 2 and 3 do not re-retrieve. Therefore differences isolate the value of added context around the same retrieved sentence.",
            ],
            [
                "Interpretation",
                "The test is qualitative until human gold alignments/classifications are available. More context is not automatically better; inspect whether it adds relevant evidence or unrelated disclosure.",
            ],
            [
                "Next step",
                "If one context approach is clearly preferable, feed that retrieved evidence to the persistent-versus-novel classifier (NLI or LLM) and benchmark against human labels.",
            ],
        ],
        columns=["item", "guidance"],
    )

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        annotation.to_excel(writer, sheet_name="Annotation", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)
        metadata.to_excel(writer, sheet_name="Metadata", index=False)
        guide.to_excel(writer, sheet_name="Guide", index=False)

    wb = load_workbook(output)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for col in range(1, ws.max_column + 1):
            letter = ws.cell(1, col).column_letter
            header = str(ws.cell(1, col).value or "")

            if any(
                key in header
                for key in ["sentence", "context", "paragraph", "notes"]
            ):
                ws.column_dimensions[letter].width = 70
            elif "cosine" in header:
                ws.column_dimensions[letter].width = 14
            elif "index" in header:
                ws.column_dimensions[letter].width = 18
            else:
                ws.column_dimensions[letter].width = 20

    wb["Annotation"].freeze_panes = "D2"
    wb["Annotation"].auto_filter.ref = wb["Annotation"].dimensions
    wb["Guide"].column_dimensions["A"].width = 28
    wb["Guide"].column_dimensions["B"].width = 110

    wb.save(output)


def main() -> None:
    ap = argparse.ArgumentParser(
        description=(
            "Compare best-sentence, +/-1-sentence, and containing-paragraph "
            "retrieval contexts on the MD&A benchmark."
        )
    )
    ap.add_argument(
        "--benchmark",
        type=Path,
        default=Path("japanese_mdna_alignment_candidates.xlsx"),
        help="Existing benchmark workbook.",
    )
    ap.add_argument(
        "--output",
        type=Path,
        default=Path("japanese_mdna_retrieval_context_test.xlsx"),
    )
    ap.add_argument("--top-k", type=int, default=3)
    ap.add_argument("--batch-size", type=int, default=32)
    args = ap.parse_args()

    if args.top_k < 1:
        raise ValueError("--top-k must be >= 1")
    if not args.benchmark.exists():
        raise FileNotFoundError(args.benchmark)

    annotation = pd.read_excel(args.benchmark, sheet_name="Annotation")
    metadata = pd.read_excel(args.benchmark, sheet_name="Metadata")

    required_annotation = {
        "annotation_id",
        "company",
        "current_sentence_index",
        "current_sentence",
    }
    missing = required_annotation - set(annotation.columns)
    if missing:
        raise ValueError(
            f"Benchmark Annotation sheet missing columns: {sorted(missing)}"
        )

    required_meta = {"company", "prior_path", "current_path"}
    missing_meta = required_meta - set(metadata.columns)
    if missing_meta:
        raise ValueError(
            f"Benchmark Metadata sheet missing columns: {sorted(missing_meta)}"
        )

    device = choose_device()
    print(f"Using device: {device}")

    output_frames = []
    company_stats = []

    for _, meta in metadata.iterrows():
        company = str(meta["company"])
        prior_path = Path(str(meta["prior_path"]))
        current_path = Path(str(meta["current_path"]))

        company_rows = annotation.loc[
            annotation["company"].astype(str) == company
        ].copy()

        if company_rows.empty:
            print(f"Skipping {company}: no benchmark rows")
            continue

        prior_raw = read_text(prior_path)
        current_raw = read_text(current_path)

        # Original benchmark-style sentence split.
        current = split_japanese_sentences(current_raw)

        # Paragraph-aware reconstruction for prior text.
        prior, prior_paragraphs, sentence_to_paragraph = (
            map_sentences_to_paragraphs(prior_raw)
        )

        # Safety check: paragraph-aware sentence reconstruction should match
        # the ordinary splitter exactly. If it does not, paragraph mapping is
        # not safe to use without revising the splitter.
        prior_plain = split_japanese_sentences(prior_raw)
        if prior != prior_plain:
            raise ValueError(
                f"{company}: paragraph-aware sentence reconstruction does not "
                "match the benchmark sentence split. Do not use Approach 3 "
                "until paragraph mapping is reconciled."
            )

        paragraph_counts = [
            sum(1 for p in sentence_to_paragraph if p == p_idx)
            for p_idx in range(len(prior_paragraphs))
        ]
        nonzero_counts = [x for x in paragraph_counts if x > 0]

        company_stats.append(
            {
                "company": company,
                "prior_sentences": len(prior),
                "prior_paragraphs": len(prior_paragraphs),
                "median_paragraph_sentences": (
                    float(np.median(nonzero_counts))
                    if nonzero_counts
                    else 0.0
                ),
                "max_paragraph_sentences": (
                    int(max(nonzero_counts))
                    if nonzero_counts
                    else 0
                ),
            }
        )

        print("\n" + "=" * 80)
        print(company)
        print("=" * 80)
        print(f"Prior sentences:   {len(prior):,}")
        print(f"Prior paragraphs:  {len(prior_paragraphs):,}")
        print(f"Current sentences: {len(current):,}")
        print(f"Benchmark rows:    {len(company_rows):,}")

        # Ensure workbook row indices still point to exact current text.
        for _, row in company_rows.iterrows():
            idx = int(row["current_sentence_index"])
            if idx < 0 or idx >= len(current):
                raise IndexError(
                    f"{company}: current_sentence_index {idx} out of range"
                )

            expected = normalize_sentence(row["current_sentence"])
            actual = normalize_sentence(current[idx])

            if expected != actual:
                raise ValueError(
                    f"{company}: benchmark/current text mismatch at index {idx}\n"
                    f"Workbook: {expected}\n"
                    f"File:     {actual}"
                )

        queries = company_rows["current_sentence"].astype(str).tolist()

        for short, model_name in MODELS.items():
            company_rows = add_model_results(
                frame=company_rows,
                short_name=short,
                model_name=model_name,
                prior_sentences=prior,
                prior_paragraphs=prior_paragraphs,
                sentence_to_paragraph=sentence_to_paragraph,
                current_sentences=queries,
                batch_size=args.batch_size,
                top_k=args.top_k,
                device=device,
            )

        output_frames.append(company_rows)

    if not output_frames:
        raise ValueError("No benchmark companies were processed.")

    out = pd.concat(output_frames, ignore_index=True)

    # Restore benchmark order.
    out["_annotation_sort"] = (
        out["annotation_id"]
        .astype(str)
        .str.extract(r"(\d+)", expand=False)
        .astype(int)
    )
    out = (
        out.sort_values("_annotation_sort")
        .drop(columns=["_annotation_sort"])
        .reset_index(drop=True)
    )

    summary = build_summary(
        annotation=out,
        company_stats=company_stats,
        device=device,
        top_k=args.top_k,
    )

    write_workbook(
        annotation=out,
        summary=summary,
        metadata=metadata,
        output=args.output,
    )

    print("\n" + "=" * 80)
    print(f"Wrote: {args.output}")
    print(f"Rows: {len(out):,}")
    print("\nSummary:")
    print(summary.to_string(index=False))


if __name__ == "__main__":
    main()
